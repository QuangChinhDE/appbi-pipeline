"""Build the Vietnamese AppBI Pipeline manual as a polished DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "README.md"
EVIDENCE = HERE / "kiem-chung.md"
TARGET = HERE / "Huong-dan-cai-dat-va-su-dung-AppBI-Pipeline.docx"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_font(run, name="Aptos", size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str) -> None:
    # Basic Markdown inline formatting sufficient for this manual.
    token = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]))
        value = match.group(0)
        if value.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", value).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_font(run, color="165E91")
        elif value.startswith("**"):
            set_font(paragraph.add_run(value[2:-2]), bold=True)
        elif value.startswith("`"):
            run = paragraph.add_run(value[1:-1])
            set_font(run, name="Cascadia Mono", size=9.2, color="17384D")
        else:
            set_font(paragraph.add_run(value[1:-1]), italic=True)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8.5, color="687681")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(23, 36, 46)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (
        ("Title", 27, 0, 16),
        ("Heading 1", 18, 18, 9),
        ("Heading 2", 13.5, 13, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "APPBI PIPELINE   |   HƯỚNG DẪN CÀI ĐẶT VÀ SỬ DỤNG"
    set_font(header.runs[0], size=8, bold=True, color="536674")
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_title_page(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(eyebrow.add_run("SỔ TAY CÀI ĐẶT VÀ VẬN HÀNH"), size=10, bold=True, color="195B89")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Hướng dẫn cài đặt và sử dụng AppBI Pipeline")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(sub.add_run("Từ lúc lấy source mới nhất đến đồng bộ dữ liệu và xử lý sự cố"), size=13, color="586674")
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    labels = ("Phiên bản đối chiếu", "Nhánh source", "Ngày biên soạn", "Múi giờ")
    values = ("3937ef9", "master", "08 tháng 10 năm 2026", "UTC+7")
    for i, (label, value) in enumerate(zip(labels, values)):
        meta.columns[0].width = Cm(5.2)
        meta.columns[1].width = Cm(8.5)
        for cell in meta.rows[i].cells:
            set_cell_border(cell)
            set_cell_margins(cell, 130, 150, 130, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(meta.rows[i].cells[0], "E8F1F7")
        p0 = meta.rows[i].cells[0].paragraphs[0]
        p1 = meta.rows[i].cells[1].paragraphs[0]
        set_font(p0.add_run(label), bold=True, size=9.5)
        set_font(p1.add_run(value), size=9.5)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(note.add_run("Ảnh trong tài liệu được chụp trực tiếp từ ứng dụng local đã khởi chạy và kiểm chứng."), size=9.5, italic=True, color="586674")
    doc.add_page_break()


def add_contents(doc: Document, headings: list[str]) -> None:
    doc.add_heading("Mục lục", level=1)
    for heading in headings:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        text = re.sub(r"^\d+\s+", "", heading)
        add_inline(p, text)
    doc.add_page_break()


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        return [rows[0]] + rows[2:], i
    return None, start


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for r_idx, values in enumerate(rows):
        prevent_row_split(table.rows[r_idx])
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_idx == 0:
                shade(cell, "21475F")
            elif r_idx % 2 == 0:
                shade(cell, "F2F6F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx < len(values):
                add_inline(p, values[c_idx])
            for run in p.runs:
                set_font(run, size=8.5, bold=(r_idx == 0), color=("FFFFFF" if r_idx == 0 else "17242E"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, code: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F0F4F7")
    p_pr.append(shd)
    run = p.add_run("\n".join(code))
    set_font(run, name="Cascadia Mono", size=8.3, color="17384D")


def add_image(doc: Document, alt: str, rel: str) -> None:
    path = HERE / rel
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.55))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", alt)


def append_markdown(doc: Document, content: str) -> None:
    lines = content.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line.startswith("```"):
            if in_code:
                add_code(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line == "<!-- DOCX_PAGE_BREAK -->":
            doc.add_page_break()
            i += 1
            continue
        table_rows, next_i = parse_table(lines, i)
        if table_rows:
            add_table(doc, table_rows)
            i = next_i
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image:
            add_image(doc, image.group(1), image.group(2))
            i += 1
            continue
        if line.startswith("# "):
            # Main title already lives on the cover.
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.right_indent = Cm(0.3)
            add_inline(p, line.lstrip("> "))
            for run in p.runs:
                run.font.color.rgb = RGBColor(70, 86, 98)
                run.italic = True
            i += 1
            continue
        p = doc.add_paragraph()
        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
            add_inline(p, line)
            for run in p.runs:
                set_font(run, size=8.5, italic=True, color="586674")
        else:
            add_inline(p, line)
        i += 1


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    main_text = SOURCE.read_text(encoding="utf-8")
    headings = re.findall(r"^##\s+(.+)$", main_text, flags=re.MULTILINE)
    add_contents(doc, headings)
    append_markdown(doc, main_text)
    evidence_text = EVIDENCE.read_text(encoding="utf-8")
    evidence_lines = evidence_text.splitlines()
    evidence_title = evidence_lines[0].removeprefix("# ")
    doc.add_heading(evidence_title, level=1)
    append_markdown(doc, "\n".join(evidence_lines[1:]))
    doc.core_properties.title = "Hướng dẫn cài đặt và sử dụng AppBI Pipeline"
    doc.core_properties.subject = "Cài đặt, vận hành và xử lý sự cố AppBI Pipeline"
    doc.core_properties.keywords = "AppBI, Pipeline, Docker, Airbyte, dbt"
    doc.core_properties.author = "AppBI Data Team"
    doc.save(TARGET)
    print(f"Created {TARGET} ({TARGET.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
